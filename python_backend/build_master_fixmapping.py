"""Tujuan: Menyusun "MASTER BARANG <X>.xlsx" bertaksonomi dari sheet "Fix Mapping" workbook lama.
Caller: `python build_master_fixmapping.py "<workbook.xlsx>" [keluaran.xlsx]`. Alat, bukan gerbang.
Dependensi: openpyxl. Side Effects: menulis satu berkas xlsx di `data/rebuild_master/`.
Main Functions: bangun, _pisah. Self-check di `__main__` bila dijalankan tanpa argumen.

KENAPA ADA. `loop_master_builder.py` menyusun master dari sheet "Form Fix", dan untuk KINO NON
FOOD hasilnya membuang 572 dari 685 barang: hanya 113 baris yang punya "Nama KLP", sedangkan
`_parse_master_barang_xlsx` MEMBUANG baris tanpa kelompok. Akibatnya master Kino di mesin ini
tidak memuat RESIK V maupun OVALE 2IN1 sama sekali — surat September 2026 pulang tanpa satu kode
pun, dan jalur Summary terlihat rusak padahal yang kurang datanya.

Sheet "Fix Mapping" justru lengkap: 606 barang, 606 punya NAMA KELOMPOK, 101 kelompok, dan
kode-kodenya cocok dengan yang sudah hidup di `promo_rule` produksi. Ia kurasi manusia, jadi
kelompoknya diambil APA ADANYA. Yang diturunkan cuma gramasi dan varian, dari nama barang.

ponytail: satu sheet, satu bentuk. Workbook yang "Fix Mapping"-nya berbeda kolom akan ditolak
dengan sebabnya, bukan ditebak.
"""
import os
import re
import sys

import openpyxl
from openpyxl import Workbook

BASE = os.path.dirname(os.path.abspath(__file__))
KELUAR_DIR = os.path.join(BASE, "data", "rebuild_master")

HEADERS = [
    "Nama Barang Principle", "Kode Barang", "Nama Barang", "Nama Pcpl",
    "kode  2 Digit                (hrf+No)", "Nama KLP",
    "kode  2 Digit                (Nomor)", "Nama Sub KLP",
    "kode  1 Digit                (Nomor)", "Nama Sub KLP2",
    "kode  1 Digit                (Nomor)5", "Nama            Aroma/             Rasa",
    "kode  2 Digit                (Nomor)2", "Nama            Gramasi atau Jumlah Pack per CTN",
    "kode  4 Digit                (Nomor)", "Nama Jenis Kemasan",
]

# "100ML", "50 ML", "15GR", "60M", "1LTR". Dipakai untuk MEMOTONG ekor nama, bukan menebak isi.
GRAM = re.compile(r"\b(\d+(?:[.,]\d+)?)\s*(KG|GR|GRAM|G|ML|LTR|L|CC|PCS|SHEET|SHEETS|M)\b", re.I)


def _pisah(nama, kelompok):
    """("KNF RESIK V MANJAKANI WHITENING 50ML X 72 BTL", "RESIK V MANJAKANI") -> ("WHITENING", "50ML").

    Varian = sisa nama setelah awalan distributor, nama kelompok, dan ekor ukuran/kemasan
    dibuang. Kalau tidak ada sisa, variannya memang KOSONG — dan itu bermakna: "RESIK V
    MANJAKANI" polos harus bisa dibedakan dari "RESIK V MANJAKANI WHITENING", persis kasus
    yang jadi aturan "All Variant".
    """
    t = " ".join(str(nama or "").upper().split())
    t = re.sub(r"^(KNF|KIF|PT\.?)\s+", "", t)
    m = GRAM.search(t)
    gramasi = f"{m.group(1)}{m.group(2).upper()}" if m else ""
    if m:
        t = t[:m.start()]
    k = " ".join(str(kelompok or "").upper().split())
    if k and t.startswith(k):
        t = t[len(k):]
    else:
        # Kelompok kurasi kadang memakai singkatan ("OVALE FACE MASK BEDAK D."); buang
        # kata-kata awalnya yang memang sama, satu per satu, bukan seluruh string.
        for kata in k.split():
            if t.startswith(kata + " ") or t == kata:
                t = t[len(kata):].lstrip()
            else:
                break
    return " ".join(t.split()).strip(" -"), gramasi


def bangun(sumber, keluaran=None):
    wb = openpyxl.load_workbook(sumber, read_only=True, data_only=True)
    if "Fix Mapping" not in wb.sheetnames:
        raise SystemExit(f"'{os.path.basename(sumber)}' tidak punya sheet 'Fix Mapping'.")
    ws = wb["Fix Mapping"]
    baris = list(ws.iter_rows(values_only=True))
    wb.close()
    kepala = [" ".join(str(c or "").upper().split()) for c in baris[0]]
    for wajib in ("KODE BARANG", "NAMA BARANG", "NAMA KELOMPOK"):
        if wajib not in kepala:
            raise SystemExit(f"Kolom '{wajib}' tidak ada di sheet 'Fix Mapping'. Kolom: {kepala[:8]}")
    ik, inm, ikl = kepala.index("KODE BARANG"), kepala.index("NAMA BARANG"), kepala.index("NAMA KELOMPOK")

    out = Workbook()
    wo = out.active
    wo.title = "Sheet1"
    wo.append(HEADERS)
    n = 0
    for r in baris[1:]:
        kode = str(r[ik] or "").strip()
        nama = " ".join(str(r[inm] or "").split())
        kel = " ".join(str(r[ikl] or "").split())
        if not (kode and nama and kel):
            continue
        varian, gramasi = _pisah(nama, kel)
        rec = [""] * len(HEADERS)
        rec[1], rec[2], rec[5], rec[11], rec[13] = kode, nama, kel, varian, gramasi
        wo.append(rec)
        n += 1

    nama_pcpl = os.path.splitext(os.path.basename(sumber))[0].upper()
    for potong in ("FIX_FORM MASTER BARANG -", "FIX FORM MASTER BARANG -", "FIX MASTER BARANG", "MASTER BARANG"):
        nama_pcpl = nama_pcpl.replace(potong, " ")
    nama_pcpl = " ".join(nama_pcpl.split()) or "TANPA NAMA"
    keluaran = keluaran or os.path.join(KELUAR_DIR, f"MASTER BARANG {nama_pcpl}.xlsx")
    os.makedirs(os.path.dirname(keluaran), exist_ok=True)
    out.save(keluaran)
    print(f"{n} barang -> {keluaran}")
    return keluaran, n


if __name__ == "__main__":
    if len(sys.argv) > 1:
        bangun(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
        raise SystemExit(0)
    kasus = [
        ("KNF RESIK V MANJAKANI WHITENING 50ML X 72 BTL", "RESIK V MANJAKANI", ("WHITENING", "50ML")),
        ("KNF RESIK V MANJAKANI 50ML X 72 BTL", "RESIK V MANJAKANI", ("", "50ML")),
        ("KNF OVALE FACIAL LOTION ANTI ACNE 100ML X 36 B", "OVALE FACIAL LOTION", ("ANTI ACNE", "100ML")),
        ("KNF RESIK V GODOKAN SIRIH 100ML X 36 BTL", "RESIK V GODOKAN", ("SIRIH", "100ML")),
        ("KNF ABSOLUTE CHAMOMILE 150ML X 36", "ABSOLUTE", ("CHAMOMILE", "150ML")),
    ]
    gagal = 0
    for nama, kel, harap in kasus:
        dapat = _pisah(nama, kel)
        if dapat != harap:
            gagal += 1
            print(f"GAGAL {nama!r} + {kel!r} -> {dapat}, harusnya {harap}")
        else:
            print(f"OK    {kel:<22} -> varian={dapat[0]!r:<14} gramasi={dapat[1]!r}")
    print("\nSEMUA LULUS" if not gagal else f"\n{gagal} GAGAL")
    raise SystemExit(1 if gagal else 0)
