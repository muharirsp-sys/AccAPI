# Tujuan     : Bandingkan berkas Laporan Harian hasil web dengan arsip LaporanAuto (Excel lama) sel demi sel.
# Caller     : Developer, mis. `python laporan_harian_bandingkan.py <folder_hasil> <folder_acuan> [--tanggal 2026-10-24]`.
# Dependensi : python-calamine (nilai), openpyxl (format angka/tanggal baris data pertama).
# Main Functions: compare_file(ours, ref) -> list selisih; main() mencetak ringkasan per berkas/sheet/kolom.
# Side Effects: Hanya membaca workbook; tidak menulis apa pun.
from __future__ import annotations

import argparse
import datetime as dt
import math
import os
import sys
from collections import defaultdict

from python_calamine import CalamineWorkbook

FORMAT_ROWS = 3  # format angka seragam per kolom; cek beberapa baris data pertama saja


def _norm(value):
    """Nilai tampak Excel: angka 15 digit signifikan, tanggal tanpa jam 00:00, kosong = None."""
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return None
        return float(f"{float(value):.15g}")
    if isinstance(value, dt.datetime):
        return value.date() if value.time() == dt.time() else value
    return value


def _sheets(path: str) -> dict:
    wb = CalamineWorkbook.from_path(path)
    return {name: wb.get_sheet_by_name(name).to_python(skip_empty_area=False) for name in wb.sheet_names}


def _formats(path: str) -> dict:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(min_row=2, max_row=1 + FORMAT_ROWS):
            rows.append([getattr(c, "number_format", None) if getattr(c, "value", None) is not None else None
                         for c in row])
        out[ws.title] = rows
    wb.close()
    return out


def _trim(rows: list) -> list:
    """Buang baris/kolom kosong di ujung (dimensi Excel bisa lebih lebar dari isi)."""
    rows = [list(r) for r in rows]
    while rows and all(_norm(v) is None for v in rows[-1]):
        rows.pop()
    width = max((max((i + 1 for i, v in enumerate(r) if _norm(v) is not None), default=0) for r in rows), default=0)
    return [r[:width] + [None] * (width - len(r)) for r in rows]


def compare_file(ours: str, ref: str, check_formats: bool = True) -> list:
    """Kembalikan daftar selisih (dict) antara dua workbook; kosong = identik."""
    diffs = []
    a, b = _sheets(ours), _sheets(ref)
    if list(a) != list(b):
        diffs.append({"sheet": "*", "kind": "sheetnames", "ours": list(a), "ref": list(b)})
    for index, (name_b, rows_b) in enumerate(b.items()):
        names_a = list(a)
        name_a = name_b if name_b in a else (names_a[index] if index < len(names_a) else None)
        if name_a is None:
            diffs.append({"sheet": name_b, "kind": "missing-sheet"})
            continue
        ra, rb = _trim(a[name_a]), _trim(rows_b)
        if len(ra) != len(rb):
            diffs.append({"sheet": name_b, "kind": "rowcount", "ours": len(ra), "ref": len(rb)})
        ha = ra[0] if ra else []
        hb = rb[0] if rb else []
        if [_norm(v) for v in ha] != [_norm(v) for v in hb]:
            diffs.append({"sheet": name_b, "kind": "header", "ours": ha, "ref": hb})
        for r in range(1, min(len(ra), len(rb))):
            row_a, row_b = ra[r], rb[r]
            for c in range(max(len(row_a), len(row_b))):
                va = _norm(row_a[c]) if c < len(row_a) else None
                vb = _norm(row_b[c]) if c < len(row_b) else None
                if va != vb:
                    col = hb[c] if c < len(hb) else c
                    diffs.append({"sheet": name_b, "kind": "cell", "row": r + 1, "col": col,
                                  "ours": row_a[c] if c < len(row_a) else None,
                                  "ref": row_b[c] if c < len(row_b) else None})
    if check_formats and not diffs:
        fa, fb = _formats(ours), _formats(ref)
        for index, (name_b, rows_b) in enumerate(fb.items()):
            rows_a = fa.get(name_b) or (list(fa.values())[index] if index < len(fa) else [])
            hb = b[name_b][0] if b.get(name_b) else []
            for r, (row_a, row_b) in enumerate(zip(rows_a, rows_b)):
                for c, (fmt_a, fmt_b) in enumerate(zip(row_a, row_b)):
                    general = {None, "General"}
                    if fmt_a != fmt_b and not (fmt_a in general and fmt_b in general):
                        diffs.append({"sheet": name_b, "kind": "format", "row": r + 2,
                                      "col": hb[c] if c < len(hb) else c, "ours": fmt_a, "ref": fmt_b})
    return diffs


def _keyword(file_name: str) -> str:
    stem = os.path.splitext(file_name)[0]
    return stem.split("_", 1)[1] if "_" in stem else stem


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("ours")
    parser.add_argument("ref")
    parser.add_argument("--tanggal", help="awalan nama berkas acuan, mis. 2026-10-24")
    parser.add_argument("--contoh", type=int, default=5, help="contoh selisih per kolom")
    args = parser.parse_args(argv)
    refs = {_keyword(n): n for n in sorted(os.listdir(args.ref))
            if n.lower().endswith(".xlsx") and (not args.tanggal or n.startswith(args.tanggal))}
    ours = {_keyword(n): n for n in sorted(os.listdir(args.ours)) if n.lower().endswith(".xlsx")}
    total = 0
    for keyword, ref_name in refs.items():
        ref_path = os.path.join(args.ref, ref_name)
        if keyword not in ours:
            print(f"[{keyword}] TIDAK ADA berkas hasil")
            total += 1
            continue
        try:
            diffs = compare_file(os.path.join(args.ours, ours[keyword]), ref_path)
        except Exception as exc:  # acuan rusak tetap dilaporkan, tidak menghentikan berkas lain
            print(f"[{keyword}] GAGAL dibaca: {exc}")
            total += 1
            continue
        total += len(diffs)
        if not diffs:
            print(f"[{keyword}] identik")
            continue
        print(f"[{keyword}] {len(diffs)} selisih")
        grouped = defaultdict(list)
        for d in diffs:
            grouped[(d["sheet"], d["kind"], d.get("col"))].append(d)
        for (sheet, kind, col), items in grouped.items():
            label = f"  {sheet} · {kind}" + (f" · {col}" if col is not None else "")
            print(f"{label}: {len(items)}")
            for d in items[:args.contoh]:
                where = f"baris {d['row']}: " if "row" in d else ""
                print(f"      {where}hasil={d.get('ours')!r} acuan={d.get('ref')!r}"[:400])
    print(f"TOTAL selisih: {total}")
    return 0 if total == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
