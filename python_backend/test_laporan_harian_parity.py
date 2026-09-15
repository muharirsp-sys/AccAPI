# Tujuan: Cek format YUDI/FIX lama, SAHAR/MSM, stok HEINZ/PIC dan rekap manager.
# Caller: Developer/CI, python python_backend/test_laporan_harian_parity.py.
# Dependensi: pipeline, manager, pandas, openpyxl.
# Main Functions: main; Side Effects: workbook sementara, tanpa DB/email.
from pathlib import Path
import tempfile
import zipfile
import xml.etree.ElementTree as ET
import openpyxl
import pandas as pd
import laporan_harian as LH
from laporan_harian_manager import build_manager_rows, write_manager_report


def main():
    lk=LH.load_lookups_json()
    assert lk.principal_to_spv['MEGA SURYA MAS, PT']=='SYAMSUL'
    assert lk.sm_map['MEGA SURYA MAS, PT']=='ADNAN'
    assert lk.stock_spv_map['FORISA NUSAPERSADA, PT']=='YUDI'
    base={'NO_NOTA':'INV-TEST','TANGGAL':pd.Timestamp('2026-09-11'),'MATA_UANG':'IDR',
          'KODE_CUST':'C-TEST','KODE_SALESMAN':'S1','SALESMAN':'MSM1_MT_TEST',
          'KODE_BARANG':'SKU','NAMA_BARANG':'ITEM','QTY':1,'HARGA':1100,'NILAI JUAL':1100,
          'DPP':1000,'JUMLAH':1110,'PRINCIPLE':'MEGA SURYA MAS, PT','JENIS PRODUK':'MSM',
          'CABANG_OR_COSTCENTER':'MSM','REM':''}
    raw=pd.DataFrame([base])
    fix=LH._prep_acc(raw,lk)
    fix['TANGGAL_DATE']=fix['TANGGAL'].dt.date
    sb=LH.build_salesbase(fix,lk)
    groups,missing=LH.resolve_report_groups(sb,['SAHAR','MSM','SYAMSUL','ADNAN'],lk)
    assert not missing
    assert all(len(g['frame'])==1 for g in groups)
    rows=build_manager_rows(sb,lk)
    assert rows[0][2:6]==['ADNAN','SYAMSUL','MSM','MEGA SURYA MAS, PT']
    assert rows[0][6:9]==[1110,1110,1]
    negative=sb.copy();negative['JUMLAH']=-500.50
    neg_rows=build_manager_rows(negative,lk)
    assert neg_rows[0][7:9]==[-501,0]
    invalid=sb.copy();invalid['JUMLAH']=None
    try:build_manager_rows(invalid,lk)
    except ValueError:pass
    else:raise AssertionError('Missing amount must not become zero')
    try:build_manager_rows(sb.drop(columns=['JUMLAH']),lk)
    except ValueError:pass
    else:raise AssertionError('Missing amount column must fail clearly')
    volume=pd.DataFrame([{'JENISPRODUK':'PURATOS','NAMA_BARANG':'ITEM 12 PCS X 1,5KG','QTY_SATUANKECIL':2},
                         {'JENISPRODUK':'PURATOS','NAMA_BARANG':'ITEM X 500GR','QTY_SATUANKECIL':-2},
                         {'JENISPRODUK':'MOTASA','NAMA_BARANG':'ITEM X 5KG','QTY_SATUANKECIL':1}])
    vf=LH.build_principal_report('YUDI',volume,LH.REPORT_COLUMNS)
    assert vf.columns[-1]=='VOLUME_KG'
    assert vf['VOLUME_KG'].tolist()==[3,-1,None]
    with tempfile.TemporaryDirectory() as temp:
        legacy = sb.copy()
        legacy['QTY_SATUANKECIL'] = 2
        legacy = legacy.rename(columns={'QTY_SATUANKECIL':'FIX QTY_SATUAN KECIL'})
        legacy['GOLONGAN'] = 'YUDI'
        legacy['JENISPRODUK'] = 'PURATOS'
        legacy['NAMA_BARANG'] = 'ITEM X 500GR'
        legacy['NPWP'] = 123456789
        stock = pd.DataFrame([{'PRINCIPAL':'HEINZ ABC INDONESIA, PT','GOLONGAN':'ZUL & ARUL',
                               'KODE_BARANG':'H1','QTY AKHIR':4}])
        heinz = sb.copy()
        heinz['GOLONGAN'] = 'ZUL & ARUL'
        heinz['PRINCIPAL'] = 'HEINZ ABC INDONESIA, PT'
        written, unmatched = LH.write_report_files(pd.concat([legacy,heinz]),temp,'2026-09-11',['YUDI','HEINZ'],lk,stock)
        assert not unmatched
        assert written[1]['stockRows'] == 1
        yw = openpyxl.load_workbook(written[0]['path'],data_only=True)
        assert yw.active['AH2'].value == 2 and yw.active['AY2'].value == 1
        assert yw.active['O2'].value == '123456789'
        assert yw.sheetnames == ['YUDI','YUDI STOCK']
        yw.close()
        snapshot=LH.write_to_format_file(legacy,temp,'2026-09-11')
        sw=openpyxl.load_workbook(snapshot['path'],data_only=True)
        snapshot_headers=[c.value for c in sw.active[1]]
        assert sw.active.cell(2,snapshot_headers.index('FIX QTY_SATUAN KECIL')+1).value==2
        sw.close()
        result=write_manager_report(sb,lk,temp,'2026-09-11')
        with zipfile.ZipFile(Path(temp,result['fileName'])) as package:
            rels=ET.fromstring(package.read('xl/_rels/workbook.xml.rels'))
            assert any(r.attrib['Type'].endswith('/styles') for r in rels)
            assert any(r.attrib['Type'].endswith('/theme') for r in rels)
            content_types=ET.fromstring(package.read('[Content_Types].xml'))
            assert {'/xl/styles.xml','/xl/theme/theme1.xml','/xl/tables/table1.xml'} <= {r.get('PartName') for r in content_types}
            for name in package.namelist():
                if name.endswith('.xml'):assert b"encoding='utf8'" not in package.read(name)[:100]
        wb=openpyxl.load_workbook(Path(temp,result['fileName']),data_only=True)
        ws=wb.active
        assert [c.value for c in ws[1]]==['NO','CONCATENATE','SM','GOLONGAN','JENISPRODUK','PRINCIPAL','Sum of JUMLAH','ROUND','PENCAPAIAN','TARGET','TARGET FIX','%']
        assert ws['I2'].value==1 and ws['G3'].value==1110
        assert ws.column_dimensions['B'].hidden and ws.column_dimensions['G'].hidden
        assert all(ws.cell(2,c).value is None for c in [10,11,12])
        assert len(ws.tables)==1
        wb.close()
        wf=openpyxl.load_workbook(Path(temp,result['fileName']),data_only=False)
        assert wf.active['H2'].value=='=ROUND(G2,0)'
        assert wf.active['I2'].value=='=TRUNC(H2/1000,0)'
        wf.close()
    print('OK: SAHAR/MSM/PIC, YUDI kg, formula dan cache Pak Fahdhar, missing/negative amounts')

if __name__=='__main__':main()
