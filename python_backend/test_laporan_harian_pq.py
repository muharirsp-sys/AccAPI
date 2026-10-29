# Tujuan: Uji paritas Laporan Harian web vs Power Query ASLI workbook lama. Fixture = potongan data nyata
#         24 Okt 2026 (nama, alamat, NPWP, salesman dipseudonimkan); acuan dibuat oleh Excel menjalankan query M
#         owner (laporan_harian_excel_pq.py), bukan oleh kode replika. Gagal bila nilai/urutan/kolom bergeser.
# Caller: npm run test:python (run_checks.py) atau python python_backend/test_laporan_harian_pq.py.
# Dependensi: laporan_harian_pq, fixtures/laporan_harian.json.gz (masukan + acuan), python-calamine.
# Main Functions: main (paritas fixture), check_export (export Accurate dibaca seperti tempelan sheet Paste),
#                 check_manager (rekap Pak Fahdhar: angka, formula, paket XLSX).
# Side Effects: Menulis workbook sementara lalu menghapusnya.
import datetime as dt
import gzip
import json
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import openpyxl
import pandas as pd
from python_calamine import CalamineWorkbook

import laporan_harian as lh
import laporan_harian_pq as pq
from laporan_harian_manager import build_manager_rows, write_manager_report

FIXTURE = Path(__file__).resolve().parent / "fixtures" / "laporan_harian.json.gz"


def norm(v):
    """Nilai tampak Excel: tanggal = angka seri, angka 15 digit, kosong/error = None."""
    if v is None or v is pq.ERR or v == "":
        return None
    if isinstance(v, dict):
        v = dt.date.fromisoformat(v["date"])  # tanggal di fixture JSON
    if isinstance(v, dt.datetime):
        v = v.date()
    if isinstance(v, dt.date):
        return float((v - dt.date(1899, 12, 30)).days)
    if isinstance(v, (int, float)):
        return float(f"{float(v):.15g}")
    return v


def table(rows):
    return [[norm(v) for v in r] for r in rows]


def same(label, ours, expected):
    a, b = table(ours), table(expected)
    assert len(a) == len(b), f"{label}: {len(a)} baris, acuan {len(b)}"
    for i, (x, y) in enumerate(zip(a, b)):
        assert x == y, f"{label} baris {i + 1}: {x} != acuan {y}"


def main():
    lk = pq.load_lookups()
    with gzip.open(FIXTURE, "rt", encoding="utf-8") as f:
        fixture = json.load(f)
    masukan = {name: [[dt.date.fromisoformat(v["date"]) if isinstance(v, dict) else v for v in r] for r in rows]
               for name, rows in fixture["masukan"].items()}
    acuan = fixture["acuan"]
    fix = pq.build_fix(masukan["Paste Acc"], masukan["Paste Lap. Retur"], lk)
    same("FIX LAP PENJ", [pq.FIX_COLUMNS] + fix, acuan["fix"])
    stock = pq.master_stock(masukan["Stock"])
    sb = pq.SalesBase(fix, lk)
    for keyword, (sales, stock_rows) in acuan["sheets"].items():
        sheets = pq.build_report(keyword, sb, fix, stock)
        same(keyword, [sheets[0][1]] + sheets[0][2], sales)
        same(f"{keyword} STOCK", [sheets[1][1]] + sheets[1][2], stock_rows)
    adnan = acuan["sheets"]["ADNAN"][0]
    for keyword, principal in (("MSM", "MEGA SURYA MAS, PT"), ("DOLPHIN", "DOLPHIN, PT")):
        sheets = pq.build_report(keyword, sb, fix, stock)
        expected = [adnan[0]] + [r for r in adnan[1:] if r[pq.SX["PRINCIPAL"]] == principal]
        same(keyword, [sheets[0][1]] + sheets[0][2], expected)
        assert [s[0] for s in sheets] == ["Sheet1", "Sheet2"]

    # Tampilan arsip: TANGGAL tanggal (angka seri pada berkas manual), TGL_JT tanggal hanya FONTERRA/URC.
    with tempfile.TemporaryDirectory() as temp:
        for keyword, serial in (("FONTERRA", False), ("MOTASA MKS 1", True), ("SOFYAN", False)):
            path = str(Path(temp, f"{keyword}.xlsx"))
            pq.write_workbook(path, pq.build_report(keyword, sb, fix, stock), serial_dates=serial)
            wb = CalamineWorkbook.from_path(path)
            row = wb.get_sheet_by_name(wb.sheet_names[0]).to_python()[1]
            assert isinstance(row[1], float if serial else dt.date), (keyword, row[1])
            assert isinstance(row[2], dt.date if keyword == "FONTERRA" else float), (keyword, row[2])
    assert pq.report_keyword(" motasa  mks 1 ") == "MOTASA MKS 1" and pq.report_keyword("HENDRIK") is None
    csv_path = Path(__file__).resolve().parents[1] / "config" / "mapping_laporan.csv"
    keywords = [line.split(",", 1)[0] for line in csv_path.read_text(encoding="utf-8").splitlines()[1:]
                if line and not line.startswith("#") and not line.startswith("Keyword,")]
    unknown = [k for k in keywords if pq.report_keyword(k) is None]
    assert not unknown, f"keyword mapping_laporan.csv tanpa laporan: {unknown}"
    check_export(masukan, fix, lk)
    check_manager(lk)
    print(f"OK: FIX {len(fix)} baris + {len(acuan['sheets']) + 2} laporan identik dengan Power Query Excel")


def check_export(masukan, fix, lk):
    """Export Accurate asli punya kolom pemisah kosong, judul berbeda (urutan sama) dan baris kaki laporan."""
    export_names = {"Kode Customer": "KODE PELANGGAN INDUK", "Market": "MARKET", "Region": "REGION",
                    "No.Batch/Spec": "Pengembalian Barang", "Keterangan": "Catatan", "Desa": "Jalan",
                    "Kecamatan": "Negara", "Kode Jenis Produk": "JENIS PRODUK"}

    def exported(rows, names, lead):
        pad = [None] * len(lead)
        out = [lead + [x for h in rows[0] for x in (names.get(h, h), "")]]
        out += [pad + [x for v in r for x in (v, None)] for r in rows[1:]]
        return out + [["Halaman 1"] + [None] * (len(out[0]) - 1)]

    retur = exported(masukan["Paste Lap. Retur"], export_names, [""])
    acc = exported(masukan["Paste Acc"], {"NO_SO": "NO. SO", None: "JENIS_TRANSAKSI", "DISKON": "Diskon %"}, [])
    rebuilt = pq.build_fix(lh._paste_rows(acc, lh.PASTE_ACC_HEADERS, "penjualan"), lh._retur_rows(retur), lk)
    same("FIX dari export Accurate", [pq.FIX_COLUMNS] + rebuilt, [pq.FIX_COLUMNS] + fix)


def check_manager(lk):
    sb = pd.DataFrame([{"PRINCIPAL": "MEGA SURYA MAS, PT", "JENISPRODUK": "MSM", "CABANG_OR_COSTCENTER": "MSM",
                        "JUMLAH": 1110}])
    rows = build_manager_rows(sb, lk)
    assert rows[0][2:6] == ["ADNAN", "SYAMSUL", "MSM", "MEGA SURYA MAS, PT"]
    assert rows[0][6:9] == [1110, 1110, 1]
    assert build_manager_rows(sb.assign(JUMLAH=-500.50), lk)[0][7:9] == [-501, 0]
    for broken in (sb.assign(JUMLAH=None), sb.drop(columns=["JUMLAH"])):
        try:
            build_manager_rows(broken, lk)
        except ValueError:
            pass
        else:
            raise AssertionError("JUMLAH kosong/tidak ada harus gagal jelas, bukan menjadi nol")
    with tempfile.TemporaryDirectory() as temp:
        result = write_manager_report(sb, lk, temp, "2026-09-11")
        with zipfile.ZipFile(Path(temp, result["fileName"])) as package:
            rels = ET.fromstring(package.read("xl/_rels/workbook.xml.rels"))
            assert any(r.attrib["Type"].endswith("/styles") for r in rels)
            assert any(r.attrib["Type"].endswith("/theme") for r in rels)
            parts = {r.get("PartName") for r in ET.fromstring(package.read("[Content_Types].xml"))}
            assert {"/xl/styles.xml", "/xl/theme/theme1.xml", "/xl/tables/table1.xml"} <= parts
        wb = openpyxl.load_workbook(Path(temp, result["fileName"]), data_only=True)
        ws = wb.active
        assert [c.value for c in ws[1]] == ["NO", "CONCATENATE", "SM", "GOLONGAN", "JENISPRODUK", "PRINCIPAL",
                                            "Sum of JUMLAH", "ROUND", "PENCAPAIAN", "TARGET", "TARGET FIX", "%"]
        assert ws["I2"].value == 1 and ws["G3"].value == 1110
        assert ws.column_dimensions["B"].hidden and ws.column_dimensions["G"].hidden
        assert len(ws.tables) == 1
        wb.close()
        wf = openpyxl.load_workbook(Path(temp, result["fileName"]), data_only=False)
        assert wf.active["H2"].value == "=ROUND(G2,0)" and wf.active["I2"].value == "=TRUNC(H2/1000,0)"
        wf.close()


if __name__ == "__main__":
    main()
